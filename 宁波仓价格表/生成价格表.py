"""从宁波仓产品目录生成可直接发给客户的价格表。"""

from __future__ import annotations

import html
import json
import math
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / "tools" / "product-data.js"
CUSTOMER_DIR = HERE / "客户版"
TITLE = "MM苗苗柔岩板价格表"
HTML_FILE = CUSTOMER_DIR / f"{TITLE}.html"
EXCEL_FILE = CUSTOMER_DIR / f"{TITLE}.xlsx"
FACTORY_UPLIFT = 10
PRINT_NOTE = "打印另加 10 元/㎡"


def read_catalog():
    program = """
        const fs = require('node:fs');
        const vm = require('node:vm');
        const context = { window: {} };
        vm.runInNewContext(fs.readFileSync(process.argv[1], 'utf8'), context);
        process.stdout.write(JSON.stringify(context.window.JieGeProductData.ningboProductCatalog));
    """
    result = subprocess.run(
        ["node", "-e", program, str(SOURCE)],
        check=True, capture_output=True, text=True,
    )
    return json.loads(result.stdout)


def extract_rows(catalog):
    rows = []
    seen = set()
    for product in catalog:
        name, source_price, specs = product["name"], product["price"], product["specs"]
        if not isinstance(name, str) or not name.strip():
            raise ValueError("产品名称缺失")
        if isinstance(source_price, bool) or not isinstance(source_price, (int, float)) or not math.isfinite(source_price):
            raise ValueError(f"价格无效：{name}")
        if not isinstance(specs, list) or not specs:
            raise ValueError(f"规格缺失：{name}")
        for spec in specs:
            if not isinstance(spec, str) or not spec.strip():
                raise ValueError(f"规格无效：{name}")
            key = (name, spec)
            if key in seen:
                raise ValueError(f"产品规格重复：{name} {spec}")
            seen.add(key)
            rows.append({"name": name, "spec": spec, "factory_price": source_price + FACTORY_UPLIFT})
    return rows


def display_price(value):
    return f"{value:g}"


def write_html(rows):
    template = (HERE / "价格表模板.html").read_text(encoding="utf-8")
    parts = []
    previous_name = ""
    for row in rows:
        name = html.escape(row["name"])
        spec = html.escape(row["spec"].replace("*", " × "))
        search = html.escape(f'{row["name"]} {row["spec"]}', quote=True)
        section = " section-start" if row["name"] != previous_name else ""
        previous_name = row["name"]
        parts.append(
            f'<tr class="price-row{section}" data-search="{search}">'
            f'<td data-label="产品名称" class="name">{name}</td>'
            f'<td data-label="规格" class="spec">{spec}<span class="unit">mm</span></td>'
            f'<td data-label="出厂价" class="price">¥{display_price(row["factory_price"])}</td>'
            f'<td data-label="备注" class="note">{PRINT_NOTE}</td>'
            '</tr>'
        )
    output = template.replace("__ROWS__", "\n".join(parts))
    output = output.replace("__ROW_COUNT__", str(len(rows)))
    HTML_FILE.write_text(output, encoding="utf-8")


def write_excel(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "柔岩板价格表"
    sheet.sheet_view.showGridLines = False
    ink, muted, blue = "1D1D1F", "6E6E73", "0066CC"
    line, soft_blue = "E5E5EA", "F5F9FF"
    sheet.merge_cells("A1:D2")
    title = sheet["A1"]
    title.value = TITLE
    title.font = Font(name="Hiragino Sans GB", size=22, bold=True, color=ink)
    title.alignment = Alignment(vertical="center", indent=1)
    title.fill = PatternFill("solid", fgColor="FFFFFF")
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[3].height = 14

    headers = ["产品名称", "规格（mm）", "出厂价（元/㎡）", "备注"]
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(4, column, label)
        cell.fill = PatternFill("solid", fgColor="F5F5F7")
        cell.font = Font(name="Hiragino Sans GB", size=10, bold=True, color=muted)
        cell.alignment = Alignment(vertical="center", horizontal="right" if column == 3 else "left", indent=1)
        cell.border = Border(bottom=Side(style="thin", color=line))
    sheet.row_dimensions[4].height = 32

    for index, row in enumerate(rows, start=1):
        position = index + 4
        values = [row["name"], row["spec"], row["factory_price"], PRINT_NOTE]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(position, column, value)
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if index % 2 else "FBFBFD")
            cell.font = Font(name="Hiragino Sans GB", size=11 if column == 3 else 10,
                             bold=column in (1, 3), color=blue if column == 3 else (muted if column == 4 else ink))
            cell.alignment = Alignment(vertical="center", horizontal="right" if column == 3 else "left", indent=1)
            cell.border = Border(bottom=Side(style="hair", color=line))
        sheet.cell(position, 3).fill = PatternFill("solid", fgColor=soft_blue)
        sheet.cell(position, 3).number_format = '"¥"#,##0.##'
        sheet.row_dimensions[position].height = 29

    for column, width in {"A": 31, "B": 23, "C": 23, "D": 27}.items():
        sheet.column_dimensions[column].width = width
    last_row = len(rows) + 4
    sheet.auto_filter.ref = f"A4:D{last_row}"
    sheet.freeze_panes = "C5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:D{last_row}"
    workbook.properties.title = TITLE
    workbook.save(EXCEL_FILE)


def verify_outputs(rows):
    workbook = load_workbook(EXCEL_FILE, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == len(rows) + 4
    assert sheet["A1"].value == TITLE
    assert [sheet.cell(4, column).value for column in range(1, 5)] == [
        "产品名称", "规格（mm）", "出厂价（元/㎡）", "备注"
    ]
    for position, row in enumerate(rows, start=5):
        assert [sheet.cell(position, column).value for column in range(1, 5)] == [
            row["name"], row["spec"], row["factory_price"], PRINT_NOTE
        ]
    workbook.close()
    content = HTML_FILE.read_text(encoding="utf-8")
    assert content.count('class="price-row') == len(rows)
    for internal_label in ("原表底价", "原表价", "开单目录", "打印参考价"):
        assert internal_label not in content


def main():
    rows = extract_rows(read_catalog())
    CUSTOMER_DIR.mkdir(exist_ok=True)
    write_html(rows)
    write_excel(rows)
    verify_outputs(rows)
    print(f"已生成客户版：{len({row['name'] for row in rows})}款产品，{len(rows)}组规格")


if __name__ == "__main__":
    main()
